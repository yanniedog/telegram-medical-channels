"""Export workbook to .xls using xlwt (sheet names truncated to 31 chars)."""
from pathlib import Path

import xlwt
from openpyxl import load_workbook

src = Path("medical_telegram_channels_comprehensive.xlsx")
dst = Path("medical_telegram_channels_comprehensive.xls")

wb_in = load_workbook(src, read_only=True, data_only=True)
wb_out = xlwt.Workbook()

for name in wb_in.sheetnames:
    ws_in = wb_in[name]
    ws_out = wb_out.add_sheet(name[:31])
    for r_idx, row in enumerate(ws_in.iter_rows(values_only=True)):
        for c_idx, val in enumerate(row):
            if val is not None:
                ws_out.write(r_idx, c_idx, val)

wb_out.save(str(dst))
print(f"Wrote {dst}")
